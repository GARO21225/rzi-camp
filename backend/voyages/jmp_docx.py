"""
JMP (Journey Management Plan) — généré À PARTIR DU FICHIER MODÈLE FOURNI
(voyages/templates_docx/jmp_modele.docx, converti tel quel depuis le .doc
d'origine) : mise en page, logos, tableaux, textes fixes, niveaux d'alerte et
règles de conduite sont ceux du modèle ; seules les valeurs variables sont
remplacées. Sortie : .docx, et PDF via LibreOffice (conversion du même fichier).
"""
import copy
import datetime
import io
import os
import re
import shutil
import subprocess
import tempfile

import docx

from accounts.models import Parametre
from .models import Rotation, Voyage, VehiculeFlotte

MODELE = os.path.join(os.path.dirname(__file__), "templates_docx", "jmp_modele.docx")


class JMPErreur(Exception):
    pass


def _cellules(row):
    vues, out = [], []
    for c in row.cells:
        if c._tc in vues:
            continue
        vues.append(c._tc)
        out.append(c)
    return out


def _set_par(par, texte):
    """Remplace le texte d'un paragraphe en gardant la mise en forme du 1er run."""
    runs = par.runs
    if not runs:
        par.add_run(texte)
        return
    runs[0].text = texte
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def _set_cell(cell, texte):
    _set_par(cell.paragraphs[0], texte)


def _set_cell_like(cell, texte, ref_run):
    """Écrit `texte` dans une cellule (vide ou non) avec la mise en forme d'un run de référence."""
    par = cell.paragraphs[0]
    for r in list(par.runs):
        r._r.getparent().remove(r._r)
    run = par.add_run(texte)
    if ref_run is not None and ref_run._r.rPr is not None:
        run._r.insert(0, copy.deepcopy(ref_run._r.rPr))
    run.font.bold = False
    return run


def _cloner_ligne(table, modele_row, apres_row=None):
    tr = copy.deepcopy(modele_row._tr)
    (apres_row or modele_row)._tr.addnext(tr)
    return docx.table._Row(tr, table)


def _fmt_date(d):
    return d.strftime("%d/%m/%Y") if d else ""


def _heure(t):
    return f"{t.hour:02d}h{t.minute:02d}" if t else ""


def _nom(p):
    return f"{p.nom} {p.prenom}".strip().upper() if p else ""


def _tel(p):
    return (getattr(p, "telephone", "") or "") if p else ""


def collecter(rotation_id):
    rot = Rotation.objects.filter(rotation_id=rotation_id).first()
    if not rot:
        raise JMPErreur("Convoi introuvable.")
    voyages = list(Voyage.objects.select_related("personnel").filter(rotation_id=rotation_id)
                   .exclude(statut="annule").exclude(statut_validation="refuse").order_by("id"))
    etapes = []
    for v in voyages:
        es = [e for e in v.etapes.all() if e.sens != "retour"]
        if es:
            etapes = sorted(es, key=lambda e: e.ordre)
            break
    param = {p.cle: p.valeur for p in Parametre.objects.filter(cle__startswith="jmp_")}
    vf = VehiculeFlotte.objects.filter(matricule=rot.vehicule_matricule).first() if rot.vehicule_matricule else None
    from residences.models import Personnel
    chauffeur = rot.conducteur_personnel or (Personnel.objects.filter(nom__iexact=rot.conducteur.split(" ")[0]).first() if rot.conducteur else None)
    second = rot.conducteur_secondaire_personnel
    return rot, voyages, etapes, param, vf, chauffeur, second


def generer_docx(rotation_id):
    rot, voyages, etapes, param, vf, chauffeur, second = collecter(rotation_id)
    d = docx.Document(MODELE)
    aujourdhui = _fmt_date(datetime.date.today())
    sat = param.get("jmp_tel_satellite") or "+8 XX"
    mtn = param.get("jmp_tel_mtn") or "05 XX"
    orange = param.get("jmp_tel_orange") or "07 XX"

    # ── En-tête (toutes les pages) : ligne d'urgence ──
    for p in d.sections[0].header.paragraphs:
        if "URGENCE" in p.text:
            runs = p.runs
            # runs : [espaces, 'URGENCE/EMERGENCY ', ': ', 'Sat Téléphone', ' : ', '+8 XX MTN / Orange : ', '07 XX/ 05 XX']
            if len(runs) >= 7:
                runs[4].text = "\xa0: "
                runs[5].text = f"{sat} MTN\xa0/ Orange : "
                runs[6].text = f"{orange}/ {mtn}"

    # ── Tableau 0 : informations générales ──
    t0 = d.tables[0]
    typ = re.sub(r"^[^A-Za-zÀ-ÿ]+", "", vf.get_categorie_display()).upper() if vf else ""
    fin = rot.date_depart if (rot.trajet_aller_seul or not rot.date_retour_prevue) else rot.date_retour_prevue
    vals = {
        (0, 1): "ROXGOLD SANGO", (0, 3): aujourdhui, (0, 5): typ,
        (1, 1): (rot.origine or "").upper(), (1, 3): (rot.destination or "").upper(),
        (1, 5): (rot.vehicule or "").upper(),
        (2, 1): _fmt_date(rot.date_depart), (2, 3): _fmt_date(fin), (2, 5): rot.vehicule_matricule or "",
    }
    for (r, c), v in vals.items():
        _set_cell(_cellules(t0.rows[r])[c], v)

    # ── Tableau 1 : trajet, équipement, manifeste ──
    t1 = d.tables[1]
    villes = ([etapes[0].origine] + [e.destination for e in etapes]) if etapes else [rot.origine, rot.destination]
    trajet = "-".join(str(v).upper() for v in villes if v)
    cell_trajet = _cellules(t1.rows[1])[0]
    _set_par(cell_trajet.paragraphs[0], trajet)
    _set_par(cell_trajet.paragraphs[1], (rot.vehicule or "").upper())
    # Numéro de téléphone satellite (ligne 4, 2e cellule)
    _set_cell(_cellules(t1.rows[4])[1], f"Numéro de téléphone satellite\xa0: {sat}")

    cc = _cellules(t1.rows[6])
    _set_cell(cc[1], _nom(chauffeur) or (rot.conducteur or "").upper())
    _set_cell(cc[3], _tel(chauffeur))
    _set_cell(cc[4], rot.origine or "")
    _set_cell(cc[5], rot.destination or "")

    cs = _cellules(t1.rows[7])
    if rot.conducteur_secondaire:
        _set_cell(cs[1], (rot.conducteur_secondaire or "").upper())
        _set_cell(cs[3], f"Chef de parcours {_tel(second)}".strip())
        _set_cell(cs[4], rot.origine or "")
        _set_cell(cs[5], rot.destination or "")
    else:
        for c in cs[1:]:
            _set_cell(c, "")

    # Lignes passagers : 1..15 dans le modèle (lignes 8..22) ; on en ajoute au besoin
    lignes = list(t1.rows[8:23])
    while len(lignes) < len(voyages):
        lignes.append(_cloner_ligne(t1, lignes[-1]))
        _set_cell(_cellules(lignes[-1])[0], str(len(lignes)))
    ref = cs[1].paragraphs[0].runs[0] if cs[1].paragraphs[0].runs else None
    for i, v in enumerate(voyages):
        c = _cellules(lignes[i])
        p = v.personnel
        for k, val in ((1, _nom(p)), (2, (getattr(p, "departement", "") or getattr(p, "societe", "") or "")),
                       (3, _tel(p)), (4, v.origine or rot.origine or ""), (5, v.destination or rot.destination or "")):
            _set_cell_like(c[k], val, ref)

    # ── Tableau 2 : côte de sécurité de route + approbation ──
    t2 = d.tables[2]
    modele_ligne = t2.rows[1]
    steps = list(t2.rows[1:7])
    if etapes:
        while len(steps) < len(etapes):
            steps.append(_cloner_ligne(t2, steps[-1]))
        for extra in steps[len(etapes):]:
            extra._tr.getparent().remove(extra._tr)
        for i, e in enumerate(etapes):
            c = _cellules(steps[i])
            _set_cell(c[1], str(e.ordre))
            _set_cell(c[2], (e.origine or "").upper())
            _set_cell(c[3], (e.destination or "").upper())
            _set_cell(c[4], f"{float(e.distance_km):g} Kms" if e.distance_km else "")
            _set_cell(c[5], _heure(e.heure_depart))
            _set_cell(c[6], _heure(e.heure_arrivee_prevue))
            _set_cell(c[7], e.pause_fatigue or "N/A")
    else:
        # aucune étape saisie : une seule ligne origine -> destination
        for extra in steps[1:]:
            extra._tr.getparent().remove(extra._tr)
        c = _cellules(steps[0])
        for k, val in ((1, "1"), (2, (rot.origine or "").upper()), (3, (rot.destination or "").upper()),
                       (4, ""), (5, _heure(rot.heure_depart) if getattr(rot, "heure_depart", None) else ""), (6, ""), (7, "N/A")):
            _set_cell(c[k], val)
    rows = t2.rows
    sig = _cellules(rows[len(rows) - 2])
    _set_cell(sig[1], _nom(chauffeur) or (rot.conducteur or "").upper())
    _set_cell(sig[5], aujourdhui)
    app = _cellules(rows[len(rows) - 1])
    _set_cell(app[1], (param.get("jmp_securite_nom") or app[1].text).upper())
    _set_cell(app[3], (param.get("jmp_securite_fonction") or app[3].text).upper())
    _set_cell(app[5], aujourdhui)

    # ── Texte : numéros d'urgence de l'étape 3 ──
    for p in d.paragraphs:
        if p.text.startswith("Téléphonez au centre d"):
            if len(p.runs) >= 2:
                p.runs[1].text = f"{orange}\xa0/ {mtn}"

    # ── Dernière page : signature du chauffeur ──
    t4 = d.tables[4]
    c4 = _cellules(t4.rows[0])
    _set_par(c4[0].paragraphs[0], f"Signature : {_nom(chauffeur) or (rot.conducteur or '').upper()}")
    runs = c4[1].paragraphs[0].runs
    if runs:
        runs[-1].text = f"Date\xa0: {aujourdhui}"

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue(), rot


def docx_vers_pdf(contenu_docx):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise JMPErreur("Conversion PDF indisponible sur ce serveur (LibreOffice absent) — téléchargez le fichier Word.")
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "jmp.docx")
        with open(src, "wb") as f:
            f.write(contenu_docx)
        profil = os.path.join(tmp, "profil")
        try:
            subprocess.run([soffice, f"-env:UserInstallation=file://{profil}", "--headless", "--convert-to", "pdf",
                            "--outdir", tmp, src], check=True, timeout=120,
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, env={**os.environ, "HOME": tmp})
        except Exception as e:  # noqa: BLE001
            raise JMPErreur(f"Échec de la conversion PDF : {e}")
        pdf = os.path.join(tmp, "jmp.pdf")
        if not os.path.exists(pdf):
            raise JMPErreur("Échec de la conversion PDF.")
        with open(pdf, "rb") as f:
            return f.read()
